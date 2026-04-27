"""Chart generation module — auto-create bar, pie, and line charts."""

from __future__ import annotations

import logging
from typing import Any

from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

# Maximum number of auto-generated charts
MAX_CHARTS = 6


def _add_bar_chart(
    ws: Worksheet,
    title: str,
    data_ws: Worksheet,
    col_idx: int,
    data_start: int,
    data_end: int,
    anchor: str,
) -> None:
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = title
    chart.y_axis.title = title
    chart.x_axis.title = None
    chart.width = 18
    chart.height = 12

    data_ref = Reference(data_ws, min_col=col_idx, min_row=data_start - 1,
                         max_row=min(data_end, data_start + 49))
    chart.add_data(data_ref, titles_from_data=True)
    chart.shape = 4
    ws.add_chart(chart, anchor)


def _add_line_chart(
    ws: Worksheet,
    title: str,
    data_ws: Worksheet,
    col_idx: int,
    data_start: int,
    data_end: int,
    anchor: str,
) -> None:
    chart = LineChart()
    chart.style = 10
    chart.title = title
    chart.y_axis.title = title
    chart.width = 18
    chart.height = 12

    data_ref = Reference(data_ws, min_col=col_idx, min_row=data_start - 1,
                         max_row=min(data_end, data_start + 99))
    chart.add_data(data_ref, titles_from_data=True)
    chart.smooth = True
    ws.add_chart(chart, anchor)


def _add_pie_chart(
    ws: Worksheet,
    title: str,
    label_range: Reference,
    data_range: Reference,
    anchor: str,
) -> None:
    chart = PieChart()
    chart.title = title
    chart.width = 16
    chart.height = 12
    chart.add_data(data_range, titles_from_data=True)
    chart.set_categories(label_range)
    ws.add_chart(chart, anchor)


def generate_charts(
    summary_ws: Worksheet,
    data_ws: Worksheet,
    col_meta: list[dict[str, Any]],
    data_start: int,
    data_end: int,
    chart_start_row: int,
) -> int:
    """Create charts on the summary sheet. Returns next available row."""
    numeric_cols = [
        m for m in col_meta if m["type"] in ("currency", "integer", "float", "percentage")
    ]
    date_cols = [m for m in col_meta if m["type"] == "date"]
    cat_cols = [m for m in col_meta if m.get("is_categorical")]

    row_cursor = chart_start_row
    chart_count = 0

    # Bar charts for numeric columns
    for meta in numeric_cols[:3]:
        if chart_count >= MAX_CHARTS:
            break
        col_letter = get_column_letter(meta["index"] + 1)
        anchor = f"A{row_cursor}"
        _add_bar_chart(
            summary_ws, meta["header"], data_ws,
            meta["index"] + 1, data_start, data_end, anchor,
        )
        chart_count += 1
        row_cursor += 16

    # Line charts for date-associated numeric columns
    if date_cols and numeric_cols:
        for meta in numeric_cols[:2]:
            if chart_count >= MAX_CHARTS:
                break
            anchor = f"A{row_cursor}"
            _add_line_chart(
                summary_ws, f"{meta['header']} Trend", data_ws,
                meta["index"] + 1, data_start, data_end, anchor,
            )
            chart_count += 1
            row_cursor += 16

    # Pie charts for categorical columns (built from pivot tables written earlier)
    # These need label+data ranges written on the summary sheet itself
    # We'll handle that in the main script when writing pivot tables.

    return row_cursor


def add_pie_from_pivot(
    ws: Worksheet,
    title: str,
    label_col: int,
    data_col: int,
    start_row: int,
    end_row: int,
    anchor: str,
) -> None:
    """Add a pie chart from a pivot table already written on the sheet."""
    labels = Reference(ws, min_col=label_col, min_row=start_row + 1, max_row=end_row)
    data = Reference(ws, min_col=data_col, min_row=start_row, max_row=end_row)
    _add_pie_chart(ws, title, labels, data, anchor)
